"""CSV/XLSX import and export."""
import csv
import datetime as dt
import io

from openpyxl import Workbook, load_workbook


def _rows_to_objects(rows: list[list]) -> list[dict]:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    header = [str(h).strip().lower() for h in rows[0]]
    return [{h: (str(r[i]).strip() if i < len(r) else "") for i, h in enumerate(header) if h} for r in rows[1:]]


def parse_csv(text: str) -> list[dict]:
    text = text.lstrip("﻿")
    first = text.split("\n", 1)[0]
    delim = ";" if first.count(";") > first.count(",") else ","
    return _rows_to_objects(list(csv.reader(io.StringIO(text), delimiter=delim)))


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.time)):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def parse_xlsx(data: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return _rows_to_objects([[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)])


def parse_upload(filename: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return parse_xlsx(data)
    if name.endswith((".csv", ".txt")):
        return parse_csv(data.decode("utf-8-sig"))
    raise ValueError("Chỉ hỗ trợ .csv hoặc .xlsx")


def to_xlsx(sheets: dict[str, list[dict]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if rows:
            cols = list(rows[0].keys())
            ws.append(cols)
            for r in rows:
                ws.append([r.get(c) for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(rows: list[dict]) -> str:
    if not rows:
        return ""
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()), lineterminator="\n")
    w.writeheader()
    w.writerows(rows)
    return "﻿" + buf.getvalue()
